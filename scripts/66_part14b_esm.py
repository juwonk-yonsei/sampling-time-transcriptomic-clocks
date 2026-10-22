#!/usr/bin/env python3
"""ESM_1.pdf and ESM_2.xlsx from stored tables. No new calculations."""

import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib import font_manager
from matplotlib.font_manager import FontProperties
from matplotlib.textpath import TextPath
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter, Transformation

ROOT = Path("/HDD8T1/WORK/GeroSci/TOD")
SUB = ROOT / "figures" / "submission" / "esm"
MANUSCRIPT = ROOT / "figures" / "manuscript"
REGULAR = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
font_manager.fontManager.addfont(REGULAR)
font_manager.fontManager.addfont(BOLD)
plt.rcParams["pdf.fonttype"] = 42
plt.rcParams["font.family"] = "Liberation Sans"

AUTHORS = "Authors: Juwon Kang, Junjeong Choi"
CORRESPONDING = ("Corresponding author: Junjeong Choi, College of Pharmacy, Yonsei Institute of Pharmaceutical "
                 "Sciences, Yonsei University, Incheon 21983, Republic of Korea; junjeong@yonsei.ac.kr")

MODULE = {
    "TM_turquoise_mortality": "TM-turquoise",
    "CMP_TAms_chrono_turquoise": "chrono-turquoise",
    "CMP_TAms_chrono_yellow": "chrono-yellow",
    "CMP_TAms_chrono_orange": "chrono-orange",
    "CMP_TAms_chrono_brown": "chrono-brown",
    "CMP_TAms_chrono_red": "chrono-red",
    "CMP_TAms_mort_turquoise": "mort-turquoise",
    "CMP_TAms_mort_darkred": "mort-darkred",
}
CLOCK = {
    "TAms_chrono": "tAge-chronological",
    "TAms_mort": "tAge-mortality",
    "RA_default": "RNAAgeCalc-DESeq2",
    "RA_Peters": "RNAAgeCalc-Peters",
    "RA_GTEx": "RNAAgeCalc-GTEx",
}
DATA = {"L1": "GSE48113", "L2": "GSE122725", "L3": "GSE295758"}
HOUR = {
    "phi_bar": "phi_bar (h)",
    "peak_h": "peak (h)",
    "circ_mean_h": "circ_mean (h)",
    "circ_sd_h": "circ_sd (h)",
    "peak_center": "peak_center (h)",
    "peak_halfwidth_95": "peak_halfwidth_95 (h)",
    "peak_GSE39445": "peak_GSE39445 (h)",
    "peak_GSE56931": "peak_GSE56931 (h)",
    "peak_GSE113883": "peak_GSE113883 (h)",
    "max_circ_dist": "max_circ_dist (h)",
    "s": "s (h)",
    "smallest_s": "smallest_s (h)",
    "W": "W (h)",
    "phi": "phi (h)",
}
VERSION = {
    "frozen": "prespecified scoring",
    "D11": "participant-mean imputation",
    "D11_participant_mean": "participant-mean imputation",
}


def load(path):
    return pd.read_csv(ROOT / path, sep="\t", dtype=str, keep_default_na=False)


def rewrite(text):
    for old, new in MODULE.items():
        text = text.replace(old, new)
    for old, new in CLOCK.items():
        text = text.replace(old, new)
    text = text.replace("preregistration", "prespecified plan")
    text = text.replace("Preregistration", "Prespecified plan")
    text = text.replace("preregistered", "prespecified")
    text = text.replace("Preregistered", "Prespecified")
    for old, new in DATA.items():
        text = re.sub(rf"(?<![A-Za-z0-9]){old}(?![A-Za-z0-9])", new, text)
    if text in VERSION:
        text = VERSION[text]
    return text


def prepare(frame, hour_columns):
    out = frame.copy()
    out.columns = [rewrite(str(c)) for c in out.columns]
    out = out.map(lambda v: rewrite(v) if isinstance(v, str) else v)
    renamed = []
    columns = []
    for name in out.columns:
        new = HOUR[name] if name in hour_columns and name in HOUR else name
        renamed.append((name, new))
        columns.append(new)
    if len(columns) != len(set(columns)):
        raise SystemExit(f"duplicate header {columns}")
    out.columns = columns
    kept = [c for c in out.columns if (out[c] != "").any()]
    return out.loc[:, kept], [pair for pair in renamed if pair[0] != pair[1] and pair[1] in kept]


def stack(parts, hour_columns):
    frames = []
    renames = []
    for source, frame in parts:
        prepared, pairs = prepare(frame, hour_columns)
        prepared.insert(0, "stored_file", source)
        frames.append(prepared)
        renames.extend(pairs)
    return pd.concat(frames, ignore_index=True, sort=False).fillna(""), renames


def write_sheet(ws, title, frame):
    ws["A1"] = title
    headers = list(frame.columns)
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(3, col, name)
        cell.number_format = "@"
    for r, row in enumerate(frame.itertuples(index=False), start=4):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, "" if value is None else str(value))
            cell.number_format = "@"
    return len(headers), 3 + len(frame)


def text_width(text, prop):
    if text == "":
        return 0.0
    if text.strip() == "":
        one = TextPath((0, 0), "n" * (len(text) + 1), prop=prop).get_extents().width
        two = TextPath((0, 0), "n" + text + "n", prop=prop).get_extents().width
        return two - one
    return TextPath((0, 0), text, prop=prop).get_extents().width


def wrap_caption(label, sentence, width):
    bold = FontProperties(fname=BOLD, size=8)
    regular = FontProperties(fname=REGULAR, size=8)
    space = text_width(" ", regular)
    label_w = text_width(label, bold)
    words = sentence.split()
    lines = []
    first = []
    used = label_w + space
    rest = words
    for i, word in enumerate(words):
        need = text_width(word, regular) + (space if first else 0)
        if used + need <= width:
            first.append(word)
            used += need
        else:
            rest = words[i:]
            break
    else:
        rest = []
    lines.append((label, " ".join(first)))
    while rest:
        line = []
        used = 0.0
        nxt = rest
        for i, word in enumerate(rest):
            need = text_width(word, regular) + (space if line else 0)
            if used + need <= width:
                line.append(word)
                used += need
            else:
                nxt = rest[i:]
                break
        else:
            nxt = []
        if not line:
            raise SystemExit(f"word wider than caption: {rest[0]}")
        lines.append(("", " ".join(line)))
        rest = nxt
    return lines


def caption_pdf(path, label, sentence, page_width):
    margin = 10
    width = page_width - 2 * margin
    lines = wrap_caption(label, sentence, width)
    line_h = 11
    band = line_h * len(lines) + 12
    fig = plt.figure(figsize=(page_width / 72, band / 72))
    fig.patch.set_facecolor("white")
    y = band - 8
    for lab, rest in lines:
        x = margin
        if lab:
            fig.text(x / page_width, y / band, lab, fontproperties=FontProperties(fname=BOLD, size=8), ha="left", va="top", color="black")
            x += text_width(lab, FontProperties(fname=BOLD, size=8)) + text_width(" ", FontProperties(fname=REGULAR, size=8))
        fig.text(x / page_width, y / band, rest, fontproperties=FontProperties(fname=REGULAR, size=8), ha="left", va="top", color="black")
        y -= line_h
    fig.savefig(path)
    plt.close(fig)
    return band


def cover_pdf(path, page_width):
    lines = [
        "Title: Time of blood sampling biases transcriptomic aging clocks: development, prespecified confirmation, and practical guidance",
        "Journal: GeroScience",
        AUTHORS,
        CORRESPONDING,
    ]
    height = 160 / 25.4 * 72
    fig = plt.figure(figsize=(page_width / 72, height / 72))
    fig.patch.set_facecolor("white")
    prop = FontProperties(fname=REGULAR, size=9)
    y = height - 28
    margin = 14
    usable = page_width - 2 * margin
    for line in lines:
        words = line.split()
        rows = []
        cur = []
        for word in words:
            trial = " ".join(cur + [word])
            if text_width(trial, prop) <= usable:
                cur.append(word)
            else:
                rows.append(" ".join(cur))
                cur = [word]
        rows.append(" ".join(cur))
        for row in rows:
            fig.text(margin / page_width, y / height, row, fontproperties=prop, ha="left", va="top")
            y -= 14
        y -= 6
    fig.savefig(path)
    plt.close(fig)


def build_pdf():
    captions = [
        ("FigS1_effects_secondary.pdf", "Fig. S1", "Effect sizes and composition mediation for the secondary clocks (RNAAgeCalc-Peters, RNAAgeCalc-GTEx), as in Fig. 3a, b and Fig. 4a"),
        ("FigS2_bias_secondary.pdf", "Fig. S2", "Bias simulation for the secondary clocks, as in Fig. 7; in d, the laboratory lines of the two clocks coincide because their biases at s = 4 h have equal size (0.19 of the total SD) and opposite sign"),
        ("FigS3_mitigation_secondary.pdf", "Fig. S3", "Mitigation for the secondary clocks, as in Fig. 8; the time correction c was defined for the primary clocks only"),
    ]
    writer = PdfWriter()
    first = PdfReader(MANUSCRIPT / captions[0][0]).pages[0]
    page_width = float(first.mediabox.width)
    cover = SUB / "_cover.pdf"
    cover_pdf(cover, page_width)
    writer.add_page(PdfReader(cover).pages[0])
    labels = ["Cover"]
    for name, label, sentence in captions:
        src = PdfReader(MANUSCRIPT / name).pages[0]
        band_path = SUB / f"_{label[-2:]}.pdf"
        band = caption_pdf(band_path, label, sentence, page_width)
        page = writer.add_blank_page(width=page_width, height=float(src.mediabox.height) + band)
        page.merge_transformed_page(src, Transformation().translate(0, band))
        page.merge_page(PdfReader(band_path).pages[0])
        labels.append(label)
        band_path.unlink()
    out = SUB / "ESM_1.pdf"
    with out.open("wb") as handle:
        writer.write(handle)
    cover.unlink()
    return labels


def build_xlsx():
    s3 = load("results/stage8/tables/TableS3.tsv")
    hour = set(HOUR)
    sheets = []
    s3a, r1 = stack([
        ("l2_sleep_start", s3.loc[s3.block == "l2_sleep_start"]),
        ("arntl_alias", s3.loc[s3.block == "arntl_alias"]),
        ("l2_h1", s3.loc[s3.block == "l2_h1"]),
        ("l2_h4", s3.loc[s3.block == "l2_h4"]),
        ("effects", s3.loc[s3.block == "effects"]),
        ("l3_prediction", s3.loc[s3.block == "l3_prediction"]),
    ], hour)
    s3b, r2 = stack([
        ("d11_modules", s3.loc[s3.block == "d11_modules"]),
        ("f3_features", s3.loc[s3.block == "f3_features"]),
    ], hour)
    mask_parts = []
    for path in (
        "results/stage7b/d1_development_audit.tsv",
        "results/stage7b/d1_development_model_m.tsv",
        "results/stage7b/d1_development_within_pearson.tsv",
    ):
        mask_parts.append((Path(path).name, load(path)))
    s3c, r3 = stack(mask_parts, hour)
    s3d, r4 = prepare(load("results/stage9/sim_parameters.tsv"), hour)
    design_parts = [(Path(path).name, load(path)) for path in (
        "results/stage9/sim_design1_bias.tsv",
        "results/stage9/sim_design1_fpr.tsv",
        "results/stage9/sim_design1_smallest_s.tsv",
        "results/stage9/sim_design1_montecarlo.tsv",
        "results/stage9/sim_design2_fpr.tsv",
        "results/stage9/sim_design3_power.tsv",
    )]
    s3e, r5 = stack(design_parts, hour)
    s3f, r6 = prepare(load("results/stage9/mitigation_summary.tsv"), hour)
    s4, r7 = stack([
        ("F1", load("prereg/frozen/F1.tsv")),
        ("F2", load("prereg/frozen/F2.tsv")),
        ("F3", load("prereg/frozen/F3.tsv")),
    ], hour)
    env, r8 = prepare(load("prereg/frozen/environment.tsv"), hour)
    manifest, r9 = prepare(load("prereg/manifest.tsv"), hour)
    compared = load("results/stage5/code_vs_zenodo_md5.tsv")
    zenodo = load("results/stage4/module_clocks/zenodo_v2_files.tsv")
    local = load("results/stage4/module_clocks/local_md5.tsv")
    archive = load("results/stage4/module_clocks/module_archive_members.tsv")
    checksum_rows = []
    missing_digest = []
    for _, row in compared.iterrows():
        z = zenodo.loc[zenodo["name"] == row.filename]
        loc = local.loc[local["zenodo_name"] == row.filename]
        arc = archive.loc[archive["name"] == row.filename]
        md5 = ""
        sha256 = ""
        klass = ""
        used = ""
        if len(z) == 1:
            md5 = z.iloc[0].md5
            klass = z.iloc[0]["class"]
            used = z.iloc[0].used_for_scoring
        elif len(loc) == 1 and loc.iloc[0].md5 != "":
            md5 = loc.iloc[0].md5
        if len(arc) == 1:
            sha256 = arc.iloc[0].sha256
        if md5 == "" and sha256 == "":
            missing_digest.append(row.filename)
        checksum_rows.append({
            "filename": row.filename,
            "md5": md5,
            "sha256": sha256,
            "code_bytes": row.code_bytes,
            "zenodo_bytes": row.zenodo_bytes,
            "md5_equal": row.md5_equal,
            "class": klass,
            "used_for_scoring": used,
        })
    if missing_digest:
        raise SystemExit("no stored digest:\n" + "\n".join(missing_digest))
    checksums, r10 = prepare(pd.DataFrame(checksum_rows), hour)
    s5, r11 = stack([
        ("environment", env),
        ("model_checksum", checksums),
        ("manifest", manifest),
    ], set())
    s6, r12 = prepare(load("results/stage5/locked_design_power.tsv"), hour)
    titles = [
        ("S1", "Table S1. Deviations D1–D12: content, timing relative to results, and reason.", *prepare(load("results/stage8/tables/TableS1.tsv"), hour)),
        ("S2", "Table S2. Sensitivity analyses: prespecified (flagged-sample exclusion, random slopes, batch-corrected counts, comprehensive annotation) and deviation-defined (complete-case, kNN and participant-mean versions; imputed composition).", *prepare(load("results/stage8/tables/TableS2.tsv"), hour)),
        ("S3a", "Table S3a. GSE122725 time-axis correction and exploratory H1 and H4", s3a, r1),
        ("S3b", "Table S3b. Module outcomes under participant-mean scoring", s3b, r2),
        ("S3c", "Table S3c. Masking experiment", s3c, r3),
        ("S3d", "Table S3d. Simulation parameters", s3d, r4),
        ("S3e", "Table S3e. Simulation results, designs 1–3", s3e, r5),
        ("S3f", "Table S3f. Mitigation", s3f, r6),
        ("S4", "Table S4. Frozen development parameters: peak times, amplitudes, variance components, expected morning−evening effects and module outcomes.", s4, r7),
        ("S5", "Table S5. Software versions, model-file checksums and the frozen manifest.", s5, r11 + r8 + r9 + r10),
        ("S6", "Table S6. Prospective power of the confirmatory tests.", s6, r12),
    ]
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "Title: Time of blood sampling biases transcriptomic aging clocks: development, prespecified confirmation, and practical guidance"
    cover["A2"] = "Journal: GeroScience"
    cover["A3"] = AUTHORS
    cover["A4"] = CORRESPONDING
    info = [("Cover", 1, 4, 1, cover["A1"].value)]
    for name, title, frame, _renames in titles:
        ws = wb.create_sheet(name)
        ncol, nrow = write_sheet(ws, title, frame)
        info.append((name, nrow, ncol, len(frame), title))
    wb.save(SUB / "ESM_2.xlsx")
    return info


def scan(path, needles):
    if path.suffix == ".pdf":
        import subprocess
        raw = subprocess.check_output(["pdftotext", str(path), "-"], text=True, errors="replace")
        return {n: raw.count(n) for n in needles}
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False)
    counts = {n: 0 for n in needles}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if not isinstance(value, str):
                    continue
                for n in needles:
                    counts[n] += value.count(n)
    return counts


def main():
    SUB.mkdir(parents=True, exist_ok=True)
    labels = build_pdf()
    info = build_xlsx()
    print("PDF_PAGES", labels)
    for row in info:
        print("SHEET", row[0], "rows", row[1], "cols", row[2], "data", row[3])
    needles = ["reregist", "TAms_", "RA_", "CMP_", "L1", "L2", "L3"]
    for name in ("ESM_1.pdf", "ESM_2.xlsx"):
        print(name, scan(SUB / name, needles))


if __name__ == "__main__":
    main()
